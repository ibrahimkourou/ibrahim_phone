# -*- coding: utf-8 -*-
"""
IBRAHIM PHONE - Logiciel de gestion de stock
CustomTkinter + Excel (openpyxl)

Installation :
    py -m pip install customtkinter openpyxl

Lancement :
    py gestion_stock.py
"""

import datetime
import os
import re
import sys
import tkinter as tk
from functools import partial
from tkinter import ttk

import customtkinter as ctk
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# 1. CONFIGURATION
# ============================================================

STORE_NAME = "IBRAHIM PHONE"
TAGLINE = "Gestion de stock"
EXCEL_FILENAME = "stock_telephones.xlsx"
CURRENCY = "DH"

COLORS = {
    "accent": "#4F46E5",
    "accent_hover": "#4338CA",
    "accent_soft": "#E0E7FF",

    "sidebar": "#0F172A",
    "sidebar_hover": "#1E293B",
    "sidebar_text": "#CBD5E1",

    "bg": "#F1F5F9",
    "card": "#FFFFFF",
    "field": "#F8FAFC",
    "border": "#E2E8F0",

    "text": "#0F172A",
    "muted": "#64748B",

    "success": "#16A34A",
    "success_soft": "#DCFCE7",

    "danger": "#DC2626",
    "danger_hover": "#B91C1C",
    "danger_soft": "#FEE2E2",
    "danger_soft_hover": "#FECACA",

    "warning": "#D97706",

    "violet": "#7C3AED",
    "violet_soft": "#EDE9FE",

    "sky": "#0284C7",
    "sky_soft": "#E0F2FE",
}

C = COLORS

DAYS = [
    "lundi", "mardi", "mercredi", "jeudi",
    "vendredi", "samedi", "dimanche"
]

MONTHS = [
    "janvier", "février", "mars", "avril",
    "mai", "juin", "juillet", "août",
    "septembre", "octobre", "novembre", "décembre"
]


# ============================================================
# 2. OUTILS
# ============================================================

def app_dir():
    """
    Retourne le dossier dans lequel se trouve le programme.
    Le fichier Excel sera enregistré dans ce dossier.
    """
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)

    return os.path.dirname(os.path.abspath(__file__))


def F(size=13, weight="normal"):
    """Police utilisée dans l'application."""
    return ctk.CTkFont(
        family="Segoe UI",
        size=size,
        weight=weight
    )


def ease_out(t):
    """Animation douce."""
    return 1 - (1 - t) ** 3


def hex_to_rgb(h):
    h = h.lstrip("#")

    return tuple(
        int(h[i:i + 2], 16)
        for i in (0, 2, 4)
    )


def mix(c1, c2, t):
    """Mélange deux couleurs."""
    a = hex_to_rgb(c1)
    b = hex_to_rgb(c2)

    return "#%02x%02x%02x" % tuple(
        int(a[i] + (b[i] - a[i]) * t)
        for i in range(3)
    )


def fmt_money(value):
    s = f"{value:,.2f}"
    s = s.replace(",", " ").replace(".", ",")

    return f"{s} {CURRENCY}"


def parse_price(text):
    """
    Transforme :
        799,99
        799.99
        799,99 €
    en nombre.
    """

    t = text.replace("\u00a0", "")
    t = t.replace(" ", "")

    for symbol in (CURRENCY, "€", "$"):
        t = t.replace(symbol, "")

    t = t.replace(",", ".")

    if not re.fullmatch(r"\d+(\.\d{1,2})?", t):
        return None

    value = float(t)

    if value <= 0 or value > 10_000_000:
        return None

    return round(value, 2)


# ============================================================
# 3. EXCEL
# ============================================================

class StockError(Exception):
    """Erreur liée au stock."""


class DuplicateSerialError(StockError):
    """Numéro de série déjà existant."""


def clean_serial(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    return str(value).strip()


class ExcelStock:

    HEADERS = [
        "Numéro de série",
        "Nom du téléphone",
        "Stockage",
        "Prix"
    ]

    SHEET = "Stock"

    def __init__(self, path):
        self.path = path

    # --------------------------------------------------------
    # Création du fichier
    # --------------------------------------------------------

    def _ensure(self):

        if os.path.exists(self.path):
            return

        try:
            wb = Workbook()

            ws = wb.active
            ws.title = self.SHEET

            ws.append(self.HEADERS)

            self._format_sheet(ws)

            wb.save(self.path)

        except PermissionError:

            raise StockError(
                "Impossible de créer le fichier Excel ici. "
                "Placez le programme dans un autre dossier."
            )

        except OSError as e:

            raise StockError(
                f"Impossible de créer le fichier Excel : {e}"
            )

    @staticmethod
    def _format_sheet(ws):

        fill = PatternFill(
            "solid",
            start_color="4F46E5",
            end_color="4F46E5"
        )

        for cell in ws[1]:

            cell.font = Font(
                bold=True,
                color="FFFFFF",
                size=12
            )

            cell.fill = fill

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        widths = [28, 34, 16, 18]

        for i, width in enumerate(widths, start=1):

            ws.column_dimensions[
                get_column_letter(i)
            ].width = width

        ws.row_dimensions[1].height = 26

        ws.freeze_panes = "A2"

    # --------------------------------------------------------
    # Ouverture
    # --------------------------------------------------------

    def _open(self):

        self._ensure()

        try:

            wb = load_workbook(self.path)

        except PermissionError:

            raise StockError(
                "Le fichier Excel est verrouillé. "
                "Fermez Excel puis réessayez."
            )

        except Exception:

            raise StockError(
                "Le fichier Excel est illisible ou corrompu."
            )

        if self.SHEET in wb.sheetnames:
            ws = wb[self.SHEET]
        else:
            ws = wb.active

        return wb, ws

    # --------------------------------------------------------
    # Sauvegarde
    # --------------------------------------------------------

    def _save(self, wb):

        try:

            wb.save(self.path)

        except PermissionError:

            raise StockError(
                "Le fichier Excel est ouvert. "
                "Fermez-le puis réessayez."
            )

        except OSError as e:

            raise StockError(
                f"Impossible d'enregistrer le fichier Excel : {e}"
            )

    # --------------------------------------------------------
    # Recherche
    # --------------------------------------------------------

    @staticmethod
    def _find_row(ws, serial):

        key = serial.strip().upper()

        for r in range(2, ws.max_row + 1):

            value = clean_serial(
                ws.cell(r, 1).value
            )

            if value.upper() == key:
                return r

        return None

    @staticmethod
    def _next_row(ws):

        r = ws.max_row

        while r > 1:

            value = clean_serial(
                ws.cell(r, 1).value
            )

            if value != "":
                break

            r -= 1

        return r + 1

    # --------------------------------------------------------
    # Écriture
    # --------------------------------------------------------

    @staticmethod
    def _write_row(ws, row, phone):

        c1 = ws.cell(
            row,
            1,
            phone["serial"]
        )

        c1.number_format = "@"

        ws.cell(
            row,
            2,
            phone["name"]
        )

        c3 = ws.cell(
            row,
            3,
            phone["storage"]
        )

        c3.alignment = Alignment(
            horizontal="center"
        )

        c4 = ws.cell(
            row,
            4,
            float(phone["price"])
        )

        c4.number_format = (
            f'#,##0.00 "{CURRENCY}"'
        )

    # --------------------------------------------------------
    # Lire tout le stock
    # --------------------------------------------------------

    def get_all(self):

        wb, ws = self._open()

        phones = []

        for r in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if not r:
                continue

            if clean_serial(r[0]) == "":
                continue

            try:

                price = (
                    float(r[3])
                    if len(r) > 3 and r[3] not in (None, "")
                    else 0.0
                )

            except (TypeError, ValueError):

                price = 0.0

            phones.append({
                "serial": clean_serial(r[0]),

                "name": (
                    str(r[1] or "").strip()
                    if len(r) > 1
                    else ""
                ),

                "storage": (
                    str(r[2] or "").strip()
                    if len(r) > 2
                    else ""
                ),

                "price": price
            })

        return phones

    # --------------------------------------------------------
    # Ajouter
    # --------------------------------------------------------

    def add(self, phone):

        wb, ws = self._open()

        if self._find_row(ws, phone["serial"]):

            raise DuplicateSerialError(
                "Numéro de série déjà existant."
            )

        row = self._next_row(ws)

        self._write_row(
            ws,
            row,
            phone
        )

        self._save(wb)

    # --------------------------------------------------------
    # Modifier
    # --------------------------------------------------------

    def update(self, old_serial, phone):

        wb, ws = self._open()

        row = self._find_row(
            ws,
            old_serial
        )

        if row is None:

            raise StockError(
                "Ce téléphone est introuvable."
            )

        if (
            phone["serial"].strip().upper()
            != old_serial.strip().upper()
        ):

            if self._find_row(
                ws,
                phone["serial"]
            ):

                raise DuplicateSerialError(
                    "Numéro de série déjà existant."
                )

        self._write_row(
            ws,
            row,
            phone
        )

        self._save(wb)

    # --------------------------------------------------------
    # Supprimer
    # --------------------------------------------------------

    def delete(self, serial):

        wb, ws = self._open()

        row = self._find_row(
            ws,
            serial
        )

        if row is None:

            raise StockError(
                "Ce téléphone est introuvable."
            )

        ws.delete_rows(row)

        self._save(wb)


# ============================================================
# 4. BOUTONS
# ============================================================

class HoverButton(ctk.CTkButton):

    def __init__(
        self,
        master,
        base,
        hover,
        **kwargs
    ):

        super().__init__(
            master,
            fg_color=base,
            hover=False,
            **kwargs
        )

        self._base = base
        self._hover = hover
        self._cur = base
        self._job = None

        self.bind(
            "<Enter>",
            lambda e: self._go(self._hover)
        )

        self.bind(
            "<Leave>",
            lambda e: self._go(self._base)
        )

    def set_colors(
        self,
        base,
        hover
    ):

        self._base = base
        self._hover = hover
        self._cur = base

        if self._job:

            self.after_cancel(
                self._job
            )

            self._job = None

        self.configure(
            fg_color=base
        )

    def _go(self, target):

        if self._job:

            self.after_cancel(
                self._job
            )

        start = self._cur
        steps = 8

        def tick(i=1):

            self._cur = mix(
                start,
                target,
                i / steps
            )

            self.configure(
                fg_color=self._cur
            )

            if i < steps:

                self._job = self.after(
                    15,
                    lambda: tick(i + 1)
                )

            else:

                self._job = None

        tick()


# ============================================================
# 5. MESSAGE
# ============================================================

class Banner(ctk.CTkFrame):

    STYLES = {
        "success": (
            C["success"],
            "✓"
        ),

        "error": (
            C["danger"],
            "✕"
        ),

        "warning": (
            C["warning"],
            "!"
        ),

        "info": (
            C["accent"],
            "i"
        )
    }

    def __init__(self, master):

        super().__init__(
            master,
            corner_radius=16,
            fg_color=C["success"]
        )

        self.icon = ctk.CTkLabel(
            self,
            text="✓",
            width=30,
            height=30,
            corner_radius=15,
            fg_color="white",
            text_color=C["success"],
            font=F(15, "bold")
        )

        self.icon.pack(
            side="left",
            padx=(14, 10),
            pady=12
        )

        self.label = ctk.CTkLabel(
            self,
            text="",
            text_color="white",
            font=F(14, "bold")
        )

        self.label.pack(
            side="left",
            padx=(0, 20),
            pady=12
        )

        self._y = -90
        self._token = 0
        self._hide_job = None
        self._visible = False

    def show(
        self,
        kind,
        text,
        duration=3800
    ):

        color, symbol = self.STYLES.get(
            kind,
            self.STYLES["info"]
        )

        self.configure(
            fg_color=color
        )

        self.icon.configure(
            text=symbol,
            text_color=color
        )

        self.label.configure(
            text=text
        )

        self._token += 1

        token = self._token

        if self._hide_job:

            self.after_cancel(
                self._hide_job
            )

            self._hide_job = None

        if not self._visible:

            self._y = -90

        self._visible = True

        self.place(
            relx=0.5,
            y=int(self._y),
            anchor="n"
        )

        self.lift()

        self._slide(
            16,
            token
        )

        self._hide_job = self.after(
            duration,
            self.hide
        )

    def hide(self):

        self._hide_job = None

        self._token += 1

        self._slide(
            -90,
            self._token,
            done=self._finish_hide
        )

    def _finish_hide(self):

        self._visible = False

        self.place_forget()

    def _slide(
        self,
        target,
        token,
        done=None,
        steps=12
    ):

        start = self._y

        def tick(i=1):

            if token != self._token:
                return

            self._y = (
                start
                + (target - start)
                * ease_out(i / steps)
            )

            self.place(
                relx=0.5,
                y=int(self._y),
                anchor="n"
            )

            if i < steps:

                self.after(
                    14,
                    lambda: tick(i + 1)
                )

            elif done:

                done()

        tick()


# ============================================================
# 6. CARTES STATISTIQUES
# ============================================================

class StatCard(ctk.CTkFrame):

    def __init__(
        self,
        master,
        title,
        icon,
        color,
        color_soft,
        formatter
    ):

        super().__init__(
            master,
            corner_radius=18,
            fg_color=C["card"],
            border_width=1,
            border_color=C["border"]
        )

        self.formatter = formatter
        self.value = 0.0
        self._job = None

        ctk.CTkLabel(
            self,
            text=icon,
            width=56,
            height=56,
            corner_radius=16,
            fg_color=color_soft,
            text_color=color,
            font=F(24)
        ).grid(
            row=0,
            column=0,
            rowspan=2,
            padx=(20, 14),
            pady=20
        )

        ctk.CTkLabel(
            self,
            text=title,
            text_color=C["muted"],
            font=F(13),
            anchor="w"
        ).grid(
            row=0,
            column=1,
            sticky="sw",
            padx=(0, 20),
            pady=(20, 0)
        )

        self.number = ctk.CTkLabel(
            self,
            text=formatter(0),
            text_color=C["text"],
            font=F(26, "bold"),
            anchor="w"
        )

        self.number.grid(
            row=1,
            column=1,
            sticky="nw",
            padx=(0, 20),
            pady=(0, 20)
        )

    def set_value(
        self,
        target,
        from_zero=True
    ):

        if self._job:

            self.after_cancel(
                self._job
            )

        if from_zero:

            self.value = 0.0

        start = self.value
        steps = 28

        def tick(i=1):

            self.value = (
                start
                + (target - start)
                * ease_out(i / steps)
            )

            self.number.configure(
                text=self.formatter(self.value)
            )

            if i < steps:

                self._job = self.after(
                    16,
                    lambda: tick(i + 1)
                )

            else:

                self.value = target

                self.number.configure(
                    text=self.formatter(target)
                )

                self._job = None

        tick()


# ============================================================
# 7. PAGE DE BASE
# ============================================================

class Page(ctk.CTkFrame):

    def __init__(
        self,
        app,
        title,
        subtitle
    ):

        super().__init__(
            app.content,
            fg_color=C["bg"],
            corner_radius=0
        )

        self.app = app

        self.header = ctk.CTkFrame(
            self,
            fg_color="transparent"
        )

        self.header.pack(
            fill="x",
            padx=32,
            pady=(26, 16)
        )

        left = ctk.CTkFrame(
            self.header,
            fg_color="transparent"
        )

        left.pack(
            side="left"
        )

        self.title_label = ctk.CTkLabel(
            left,
            text=title,
            font=F(26, "bold"),
            text_color=C["text"]
        )

        self.title_label.pack(
            anchor="w"
        )

        self.sub_label = ctk.CTkLabel(
            left,
            text=subtitle,
            font=F(13),
            text_color=C["muted"]
        )

        self.sub_label.pack(
            anchor="w"
        )

        self.body = ctk.CTkFrame(
            self,
            fg_color="transparent"
        )

        self.body.pack(
            fill="both",
            expand=True,
            padx=32,
            pady=(0, 26)
        )

    def on_show(self):
        pass


# ============================================================
# 8. DASHBOARD
# ============================================================

class DashboardPage(Page):

    def __init__(self, app):

        super().__init__(
            app,
            "Tableau de bord",
            "Vue d'ensemble de votre stock"
        )

        self._token = 0

        HoverButton(
            self.header,
            C["accent"],
            C["accent_hover"],
            text="＋   Ajouter un téléphone",
            height=44,
            corner_radius=12,
            font=F(14, "bold"),
            text_color="white",
            command=lambda: app.navigate("add")
        ).pack(
            side="right"
        )

        cards = ctk.CTkFrame(
            self.body,
            fg_color="transparent"
        )

        cards.pack(
            fill="x"
        )

        for i in range(3):

            cards.grid_columnconfigure(
                i,
                weight=1,
                uniform="cards"
            )

        self.card_count = StatCard(
            cards,
            "Téléphones en stock",
            "📱",
            C["accent"],
            C["accent_soft"],
            lambda v: f"{int(round(v))}"
        )

        self.card_value = StatCard(
            cards,
            "Valeur totale du stock",
            "💰",
            C["success"],
            C["success_soft"],
            fmt_money
        )

        self.card_avg = StatCard(
            cards,
            "Prix moyen",
            "📈",
            C["violet"],
            C["violet_soft"],
            fmt_money
        )

        self.card_count.grid(
            row=0,
            column=0,
            sticky="ew",
            padx=(0, 10)
        )

        self.card_value.grid(
            row=0,
            column=1,
            sticky="ew",
            padx=10
        )

        self.card_avg.grid(
            row=0,
            column=2,
            sticky="ew",
            padx=(10, 0)
        )

        recent = ctk.CTkFrame(
            self.body,
            corner_radius=18,
            fg_color=C["card"],
            border_width=1,
            border_color=C["border"]
        )

        recent.pack(
            fill="both",
            expand=True,
            pady=(20, 0)
        )

        ctk.CTkLabel(
            recent,
            text="Derniers téléphones ajoutés",
            font=F(16, "bold"),
            text_color=C["text"]
        ).pack(
            anchor="w",
            padx=24,
            pady=(18, 6)
        )

        self.list = ctk.CTkFrame(
            recent,
            fg_color="transparent"
        )

        self.list.pack(
            fill="both",
            expand=True,
            padx=20,
            pady=(0, 14)
        )

    def on_show(self):

        phones = self.app.fetch()

        total = sum(
            p["price"]
            for p in phones
        )

        self.card_count.set_value(
            len(phones)
        )

        self.card_value.set_value(
            total
        )

        self.card_avg.set_value(
            total / len(phones)
            if phones
            else 0
        )

        self._token += 1
        token = self._token

        for child in self.list.winfo_children():
            child.destroy()

        if not phones:

            ctk.CTkLabel(
                self.list,
                text=(
                    "Aucun téléphone pour le moment.\n"
                    "Cliquez sur « Ajouter un téléphone » "
                    "pour commencer."
                ),
                font=F(14),
                text_color=C["muted"]
            ).pack(
                expand=True,
                pady=30
            )

            return

        for i, phone in enumerate(
            reversed(phones[-5:])
        ):

            self.after(
                90 * i + 120,
                partial(
                    self._add_recent,
                    phone,
                    token
                )
            )

    def _add_recent(
        self,
        p,
        token
    ):

        if (
            token != self._token
            or not self.list.winfo_exists()
        ):
            return

        row = ctk.CTkFrame(
            self.list,
            corner_radius=12,
            fg_color=C["field"]
        )

        row.pack(
            fill="x",
            pady=4
        )

        row.grid_columnconfigure(
            1,
            weight=1
        )

        ctk.CTkLabel(
            row,
            text="📱",
            width=38,
            height=38,
            corner_radius=10,
            fg_color=C["accent_soft"],
            text_color=C["accent"],
            font=F(16)
        ).grid(
            row=0,
            column=0,
            padx=(12, 10),
            pady=8
        )

        info = ctk.CTkFrame(
            row,
            fg_color="transparent"
        )

        info.grid(
            row=0,
            column=1,
            sticky="w"
        )

        ctk.CTkLabel(
            info,
            text=p["name"],
            font=F(14, "bold"),
            text_color=C["text"],
            anchor="w"
        ).pack(
            anchor="w"
        )

        ctk.CTkLabel(
            info,
            text=f"N° série : {p['serial']}",
            font=F(12),
            text_color=C["muted"],
            anchor="w"
        ).pack(
            anchor="w"
        )

        ctk.CTkLabel(
            row,
            text=p["storage"],
            font=F(12, "bold"),
            text_color=C["accent"],
            fg_color=C["accent_soft"],
            corner_radius=8,
            width=70,
            height=26
        ).grid(
            row=0,
            column=2,
            padx=10
        )

        ctk.CTkLabel(
            row,
            text=fmt_money(p["price"]),
            font=F(14, "bold"),
            text_color=C["text"]
        ).grid(
            row=0,
            column=3,
            padx=(10, 18)
        )


# ============================================================
# 9. AJOUTER / MODIFIER
# ============================================================

class AddPage(Page):

    STORAGES = [
        "16 GB",
        "32 GB",
        "64 GB",
        "128 GB",
        "256 GB",
        "512 GB",
        "1 TB"
    ]

    def __init__(self, app):

        super().__init__(
            app,
            "Ajouter un téléphone",
            "Renseignez les informations puis enregistrez dans Excel"
        )

        self.edit_serial = None
        self.fields = {}
        self._flash_job = None

        self.card = ctk.CTkFrame(
            self.body,
            corner_radius=18,
            fg_color=C["card"],
            border_width=2,
            border_color=C["border"]
        )

        self.card.pack(
            fill="x",
            anchor="n"
        )

        form = ctk.CTkFrame(
            self.card,
            fg_color="transparent"
        )

        form.pack(
            fill="x",
            padx=18,
            pady=18
        )

        form.grid_columnconfigure(
            (0, 1),
            weight=1,
            uniform="f"
        )

        self.mode_tag = ctk.CTkLabel(
            form,
            text="",
            font=F(13, "bold"),
            text_color=C["warning"],
            fg_color="#FEF3C7",
            corner_radius=10,
            height=34,
            anchor="w"
        )

        self.mode_tag.grid(
            row=0,
            column=0,
            columnspan=2,
            sticky="ew",
            padx=12,
            pady=(0, 6)
        )

        self.mode_tag.grid_remove()

        self.e_serial = self._field(
            form,
            "serial",
            1,
            0,
            "Numéro de série",
            "ex : 356938035643809"
        )

        self.e_name = self._field(
            form,
            "name",
            1,
            1,
            "Nom du téléphone",
            "ex : iPhone 15 Pro Max"
        )

        self.e_storage = self._field(
            form,
            "storage",
            2,
            0,
            "Stockage",
            kind="combo"
        )

        self.e_price = self._field(
            form,
            "price",
            2,
            1,
            f"Prix ({CURRENCY})",
            "ex : 1199,99"
        )

        options = ctk.CTkFrame(
            form,
            fg_color="transparent"
        )

        options.grid(
            row=3,
            column=0,
            columnspan=2,
            sticky="w",
            padx=12,
            pady=(10, 4)
        )

        self.clear_after = ctk.CTkSwitch(
            options,
            text="Vider les champs après l'ajout",
            progress_color=C["accent"],
            font=F(13),
            text_color=C["text"]
        )

        self.clear_after.select()

        self.clear_after.pack(
            anchor="w"
        )

        buttons = ctk.CTkFrame(
            form,
            fg_color="transparent"
        )

        buttons.grid(
            row=4,
            column=0,
            columnspan=2,
            sticky="w",
            padx=12,
            pady=(14, 4)
        )

        self.btn_primary = HoverButton(
            buttons,
            C["accent"],
            C["accent_hover"],
            text="",
            width=240,
            height=48,
            corner_radius=12,
            font=F(15, "bold"),
            text_color="white",
            command=self.submit
        )

        self.btn_primary.pack(
            side="left"
        )

        self.btn_secondary = HoverButton(
            buttons,
            C["card"],
            C["field"],
            text="Vider les champs",
            width=170,
            height=48,
            corner_radius=12,
            font=F(14, "bold"),
            text_color=C["muted"],
            border_width=1,
            border_color=C["border"],
            command=self.on_secondary
        )

        self.btn_secondary.pack(
            side="left",
            padx=(12, 0)
        )

        ctk.CTkLabel(
            self.body,
            text=f"📄  Les données sont enregistrées dans : {app.stock.path}",
            font=F(12),
            text_color=C["muted"],
            anchor="w",
            justify="left"
        ).pack(
            anchor="w",
            pady=(14, 0)
        )

        self.set_add_mode()

    # --------------------------------------------------------
    # Champ
    # --------------------------------------------------------

    def _field(
        self,
        parent,
        key,
        row,
        col,
        label,
        placeholder=None,
        kind="entry"
    ):

        box = ctk.CTkFrame(
            parent,
            fg_color="transparent"
        )

        box.grid(
            row=row,
            column=col,
            sticky="ew",
            padx=12,
            pady=(8, 2)
        )

        ctk.CTkLabel(
            box,
            text=label,
            font=F(13, "bold"),
            text_color=C["text"]
        ).pack(
            anchor="w"
        )

        if kind == "entry":

            widget = ctk.CTkEntry(
                box,
                height=46,
                corner_radius=10,
                border_width=2,
                border_color=C["border"],
                fg_color=C["field"],
                text_color=C["text"],
                placeholder_text=placeholder,
                font=F(14)
            )

        else:

            widget = ctk.CTkComboBox(
                box,
                values=self.STORAGES,
                height=46,
                corner_radius=10,
                border_width=2,
                border_color=C["border"],
                fg_color=C["field"],
                button_color=C["accent"],
                button_hover_color=C["accent_hover"],
                text_color=C["text"],
                font=F(14),
                dropdown_font=F(14)
            )

            widget.set("")

        widget.pack(
            fill="x",
            pady=(6, 0)
        )

        error = ctk.CTkLabel(
            box,
            text="",
            text_color=C["danger"],
            font=F(12),
            anchor="w",
            height=20
        )

        error.pack(
            fill="x",
            pady=(2, 0)
        )

        widget.bind(
            "<KeyRelease>",
            partial(
                self._on_key,
                key
            )
        )

        widget.bind(
            "<Return>",
            lambda e: self.submit()
        )

        self.fields[key] = (
            widget,
            error
        )

        return widget

    # --------------------------------------------------------
    # Erreurs
    # --------------------------------------------------------

    def _on_key(
        self,
        key,
        event
    ):

        if event.keysym in (
            "Return",
            "KP_Enter",
            "Tab",
            "ISO_Left_Tab",
            "Shift_L",
            "Shift_R"
        ):
            return

        self.clear_error(key)

    def set_error(
        self,
        key,
        message
    ):

        widget, label = self.fields[key]

        widget.configure(
            border_color=C["danger"]
        )

        label.configure(
            text=message
        )

    def clear_error(self, key):

        widget, label = self.fields[key]

        widget.configure(
            border_color=C["border"]
        )

        label.configure(
            text=""
        )

    def clear_all_errors(self):

        for key in self.fields:
            self.clear_error(key)

    # --------------------------------------------------------
    # Modes
    # --------------------------------------------------------

    def _primary_text(self):

        if self.edit_serial:

            return "✎   Enregistrer les modifications"

        return "＋   Ajouter au stock"

    def clear_fields(self):

        self.e_serial.delete(
            0,
            "end"
        )

        self.e_name.delete(
            0,
            "end"
        )

        self.e_price.delete(
            0,
            "end"
        )

        self.e_storage.set("")

        self.clear_all_errors()

    def set_add_mode(self):

        self.edit_serial = None

        self.title_label.configure(
            text="Ajouter un téléphone"
        )

        self.sub_label.configure(
            text=(
                "Renseignez les informations "
                "puis enregistrez dans Excel"
            )
        )

        self.mode_tag.grid_remove()

        self.clear_after.configure(
            state="normal"
        )

        self.btn_primary.configure(
            text=self._primary_text()
        )

        self.btn_secondary.configure(
            text="Vider les champs"
        )

        self.clear_fields()

    def load_for_edit(self, phone):

        self.clear_fields()

        self.edit_serial = phone["serial"]

        self.title_label.configure(
            text="Modifier un téléphone"
        )

        self.sub_label.configure(
            text="Changez les informations puis enregistrez"
        )

        self.mode_tag.configure(
            text=(
                f"   ✎  Mode modification  •  "
                f"N° de série : {phone['serial']}"
            )
        )

        self.mode_tag.grid()

        self.clear_after.configure(
            state="disabled"
        )

        self.btn_primary.configure(
            text=self._primary_text()
        )

        self.btn_secondary.configure(
            text="Annuler"
        )

        self.e_serial.insert(
            0,
            phone["serial"]
        )

        self.e_name.insert(
            0,
            phone["name"]
        )

        self.e_storage.set(
            phone["storage"]
        )

        price = (
            f"{phone['price']:.2f}"
            .replace(".", ",")
        )

        self.e_price.insert(
            0,
            price
        )

    def on_secondary(self):

        if self.edit_serial:

            self.set_add_mode()

            self.app.navigate(
                "stock"
            )

        else:

            self.clear_fields()

            self.e_serial.focus_set()

            self.app.notify(
                "info",
                "Les champs ont été vidés.",
                1800
            )

    def on_show(self):

        self.after(
            250,
            self.e_serial.focus_set
        )

    # --------------------------------------------------------
    # Animations
    # --------------------------------------------------------

    def flash_card(self, color):

        if self._flash_job:

            self.after_cancel(
                self._flash_job
            )

        self.card.configure(
            border_color=color
        )

        self._flash_job = self.after(
            900,
            lambda: self.card.configure(
                border_color=C["border"]
            )
        )

    def flash_button(self, text):

        self.btn_primary.set_colors(
            C["success"],
            C["success"]
        )

        self.btn_primary.configure(
            text=text
        )

        def restore():

            self.btn_primary.set_colors(
                C["accent"],
                C["accent_hover"]
            )

            self.btn_primary.configure(
                text=self._primary_text()
            )

        self.after(
            1400,
            restore
        )

    # --------------------------------------------------------
    # Enregistrement
    # --------------------------------------------------------

    def submit(self):

        self.clear_all_errors()

        serial = self.e_serial.get().strip()
        name = self.e_name.get().strip()
        storage = self.e_storage.get().strip()
        price_text = self.e_price.get().strip()

        valid = True

        # Numéro de série
        if not serial:

            self.set_error(
                "serial",
                "Le numéro de série est obligatoire."
            )

            valid = False

        elif len(serial) > 40:

            self.set_error(
                "serial",
                "40 caractères maximum."
            )

            valid = False

        # Nom
        if not name:

            self.set_error(
                "name",
                "Le nom du téléphone est obligatoire."
            )

            valid = False

        # Stockage
        if not storage:

            self.set_error(
                "storage",
                "Choisissez un stockage."
            )

            valid = False

        # Prix
        price = (
            parse_price(price_text)
            if price_text
            else None
        )

        if not price_text:

            self.set_error(
                "price",
                "Le prix est obligatoire."
            )

            valid = False

        elif price is None:

            self.set_error(
                "price",
                "Prix invalide. Exemple : 799,99"
            )

            valid = False

        if not valid:

            self.flash_card(
                C["danger"]
            )

            self.app.notify(
                "warning",
                "Veuillez corriger les champs en rouge."
            )

            return

        phone = {
            "serial": serial,
            "name": name,
            "storage": storage,
            "price": price
        }

        editing = (
            self.edit_serial is not None
        )

        try:

            if editing:

                self.app.stock.update(
                    self.edit_serial,
                    phone
                )

            else:

                self.app.stock.add(
                    phone
                )

        except DuplicateSerialError:

            self.set_error(
                "serial",
                "Ce numéro existe déjà."
            )

            self.flash_card(
                C["danger"]
            )

            self.app.notify(
                "error",
                f"Le numéro « {serial} » existe déjà !",
                4800
            )

            return

        except StockError as e:

            self.app.notify(
                "error",
                str(e),
                6000
            )

            return

        self.app.pages["stock"].highlight = serial

        if editing:

            self.app.notify(
                "success",
                f"« {name} » a été modifié avec succès."
            )

            self.set_add_mode()

            self.app.navigate(
                "stock"
            )

            return

        self.app.notify(
            "success",
            f"« {name} » a été ajouté au stock."
        )

        self.flash_card(
            C["success"]
        )

        self.flash_button(
            "✓   Ajouté !"
        )

        if self.clear_after.get():

            self.clear_fields()

        self.e_serial.focus_set()


# ============================================================
# 10. PAGE STOCK
# ============================================================

class StockPage(Page):

    COLUMNS = (
        (
            "serial",
            "Numéro de série",
            260,
            "w"
        ),

        (
            "name",
            "Nom du téléphone",
            340,
            "w"
        ),

        (
            "storage",
            "Stockage",
            130,
            "center"
        ),

        (
            "price",
            "Prix",
            170,
            "e"
        ),
    )

    def __init__(self, app):

        super().__init__(
            app,
            "Stock",
            "Tous vos téléphones enregistrés dans Excel"
        )

        self.phones = []
        self.row_map = {}

        self.sort_key = None
        self.sort_reverse = False

        self.highlight = None

        self.armed = False
        self._arm_job = None

        # ----------------------------------------------------
        # Barre d'outils
        # ----------------------------------------------------

        toolbar = ctk.CTkFrame(
            self.body,
            fg_color="transparent"
        )

        toolbar.pack(
            fill="x",
            pady=(0, 14)
        )

        self.search = ctk.CTkEntry(
            toolbar,
            placeholder_text=(
                "🔍  Rechercher par numéro "
                "de série ou nom…"
            ),
            width=380,
            height=44,
            corner_radius=12,
            border_width=2,
            border_color=C["border"],
            fg_color=C["card"],
            text_color=C["text"],
            font=F(14)
        )

        self.search.pack(
            side="left"
        )

        self.search.bind(
            "<KeyRelease>",
            lambda e: self.render()
        )

        HoverButton(
            toolbar,
            C["card"],
            C["field"],
            text="⟳  Actualiser",
            width=130,
            height=44,
            corner_radius=12,
            font=F(14, "bold"),
            text_color=C["text"],
            border_width=1,
            border_color=C["border"],
            command=self.on_refresh
        ).pack(
            side="right"
        )

        self.btn_delete = HoverButton(
            toolbar,
            C["danger_soft"],
            C["danger_soft_hover"],
            text="🗑  Supprimer",
            width=140,
            height=44,
            corner_radius=12,
            font=F(14, "bold"),
            text_color=C["danger"],
            command=self.delete_clicked
        )

        self.btn_delete.pack(
            side="right",
            padx=(0, 10)
        )

        HoverButton(
            toolbar,
            C["card"],
            C["field"],
            text="✎  Modifier",
            width=130,
            height=44,
            corner_radius=12,
            font=F(14, "bold"),
            text_color=C["text"],
            border_width=1,
            border_color=C["border"],
            command=self.edit_selected
        ).pack(
            side="right",
            padx=(0, 10)
        )

        HoverButton(
            toolbar,
            C["accent"],
            C["accent_hover"],
            text="＋  Ajouter",
            width=130,
            height=44,
            corner_radius=12,
            font=F(14, "bold"),
            text_color="white",
            command=lambda: app.navigate("add")
        ).pack(
            side="right",
            padx=(0, 10)
        )

        # ----------------------------------------------------
        # Tableau
        # ----------------------------------------------------

        card = ctk.CTkFrame(
            self.body,
            corner_radius=18,
            fg_color=C["card"],
            border_width=1,
            border_color=C["border"]
        )

        card.pack(
            fill="both",
            expand=True
        )

        wrap = ctk.CTkFrame(
            card,
            fg_color="transparent"
        )

        wrap.pack(
            fill="both",
            expand=True,
            padx=14,
            pady=(14, 4)
        )

        self._setup_style()

        keys = [
            c[0]
            for c in self.COLUMNS
        ]

        self.tree = ttk.Treeview(
            wrap,
            columns=keys,
            show="headings",
            style="Stock.Treeview",
            selectmode="browse"
        )

        for key, label, width, anchor in self.COLUMNS:

            self.tree.heading(
                key,
                text=label,
                anchor=anchor,
                command=partial(
                    self.sort_by,
                    key
                )
            )

            self.tree.column(
                key,
                width=width,
                minwidth=90,
                anchor=anchor,
                stretch=True
            )

        scrollbar = ctk.CTkScrollbar(
            wrap,
            command=self.tree.yview
        )

        self.tree.configure(
            yscrollcommand=scrollbar.set
        )

        scrollbar.pack(
            side="right",
            fill="y",
            padx=(6, 0)
        )

        self.tree.pack(
            side="left",
            fill="both",
            expand=True
        )

        self.tree.tag_configure(
            "even",
            background="#FFFFFF"
        )

        self.tree.tag_configure(
            "odd",
            background="#F8FAFC"
        )

        self.tree.tag_configure(
            "new",
            background=C["success_soft"]
        )

        self.tree.bind(
            "<<TreeviewSelect>>",
            lambda e: self.disarm()
        )

        self.tree.bind(
            "<Double-1>",
            self._on_double_click
        )

        self.empty_label = ctk.CTkLabel(
            wrap,
            text="",
            font=F(14),
            text_color=C["muted"],
            justify="center"
        )

        self.footer = ctk.CTkLabel(
            card,
            text="",
            font=F(12),
            text_color=C["muted"],
            anchor="w"
        )

        self.footer.pack(
            fill="x",
            padx=22,
            pady=(2, 12)
        )

    # --------------------------------------------------------
    # Style tableau
    # --------------------------------------------------------

    def _setup_style(self):

        scale = max(
            1.0,
            self.winfo_fpixels("1i") / 96
        )

        style = ttk.Style()

        try:

            style.theme_use(
                "clam"
            )

        except tk.TclError:

            pass

        style.configure(
            "Stock.Treeview",
            background="#FFFFFF",
            fieldbackground="#FFFFFF",
            foreground=C["text"],
            rowheight=int(38 * scale),
            borderwidth=0,
            font=("Segoe UI", 11)
        )

        style.configure(
            "Stock.Treeview.Heading",
            background="#F1F5F9",
            foreground=C["muted"],
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            padding=(10, 10),
            borderwidth=0
        )

        style.map(
            "Stock.Treeview",
            background=[
                ("selected", C["accent"])
            ],
            foreground=[
                ("selected", "#FFFFFF")
            ]
        )

        style.map(
            "Stock.Treeview.Heading",
            background=[
                ("active", "#E2E8F0")
            ]
        )

        style.layout(
            "Stock.Treeview",
            [
                (
                    "Treeview.treearea",
                    {"sticky": "nswe"}
                )
            ]
        )

    # --------------------------------------------------------
    # Données
    # --------------------------------------------------------

    def on_show(self):

        self.refresh()

        if self.highlight:

            self.after(
                2800,
                self._clear_highlight
            )

    def _clear_highlight(self):

        self.highlight = None

        self.render()

    def refresh(self):

        self.phones = self.app.fetch()

        self.disarm()

        self.render()

    def on_refresh(self):

        self.refresh()

        self.app.notify(
            "info",
            "Liste actualisée depuis Excel.",
            2000
        )

    def selected_phone(self):

        selection = self.tree.selection()

        if not selection:
            return None

        return self.row_map.get(
            selection[0]
        )

    # --------------------------------------------------------
    # Tri
    # --------------------------------------------------------

    def sort_by(self, key):

        if self.sort_key == key:

            self.sort_reverse = (
                not self.sort_reverse
            )

        else:

            self.sort_key = key
            self.sort_reverse = False

        self.render()

    # --------------------------------------------------------
    # Affichage
    # --------------------------------------------------------

    def render(self):

        query = (
            self.search.get()
            .strip()
            .lower()
        )

        rows = [
            p
            for p in self.phones
            if (
                query in p["serial"].lower()
                or query in p["name"].lower()
            )
        ]

        if self.sort_key:

            k = self.sort_key

            rows.sort(
                key=lambda p:
                    p[k]
                    if k == "price"
                    else p[k].lower(),

                reverse=self.sort_reverse
            )

        current = self.selected_phone()

        current_serial = (
            current["serial"]
            if current
            else None
        )

        # En-têtes
        for key, label, _w, anchor in self.COLUMNS:

            arrow = ""

            if key == self.sort_key:

                arrow = (
                    "  ▼"
                    if self.sort_reverse
                    else "  ▲"
                )

            self.tree.heading(
                key,
                text=label + arrow,
                anchor=anchor
            )

        self.tree.delete(
            *self.tree.get_children()
        )

        self.row_map = {}

        reselect = None
        highlight_iid = None

        for i, p in enumerate(rows):

            iid = str(i)

            self.row_map[iid] = p

            tag = (
                "odd"
                if i % 2
                else "even"
            )

            if p["serial"] == self.highlight:

                tag = "new"

                highlight_iid = iid

            self.tree.insert(
                "",
                "end",
                iid=iid,
                tags=(tag,),
                values=(
                    p["serial"],
                    p["name"],
                    p["storage"],
                    fmt_money(p["price"])
                )
            )

            if p["serial"] == current_serial:

                reselect = iid

        if reselect is not None:

            self.tree.selection_set(
                reselect
            )

        if highlight_iid is not None:

            self.tree.see(
                highlight_iid
            )

        if rows:

            self.empty_label.place_forget()

        else:

            if self.phones:

                text = (
                    "Aucun résultat "
                    "pour cette recherche."
                )

            else:

                text = (
                    "Aucun téléphone dans le stock.\n"
                    "Cliquez sur « Ajouter » "
                    "pour commencer."
                )

            self.empty_label.configure(
                text=text
            )

            self.empty_label.place(
                relx=0.5,
                rely=0.5,
                anchor="center"
            )

            self.empty_label.lift()

        total = sum(
            p["price"]
            for p in rows
        )

        self.footer.configure(
            text=(
                f"{len(rows)} téléphone(s) affiché(s) "
                f"sur {len(self.phones)}   •   "
                f"Valeur affichée : {fmt_money(total)}   •   "
                f"Cliquez sur un en-tête pour trier"
            )
        )

    # --------------------------------------------------------
    # Actions
    # --------------------------------------------------------

    def _on_double_click(self, event):

        if self.tree.identify_region(
            event.x,
            event.y
        ) == "cell":

            self.edit_selected()

    def edit_selected(self):

        phone = self.selected_phone()

        if not phone:

            self.app.notify(
                "warning",
                "Sélectionnez d'abord un téléphone."
            )

            return

        self.app.start_edit(
            phone
        )

    def delete_clicked(self):

        phone = self.selected_phone()

        if not phone:

            self.app.notify(
                "warning",
                "Sélectionnez d'abord un téléphone."
            )

            return

        # Première étape
        if not self.armed:

            self.armed = True

            self.btn_delete.configure(
                text="⚠  Confirmer ?",
                text_color="white"
            )

            self.btn_delete.set_colors(
                C["danger"],
                C["danger_hover"]
            )

            self._arm_job = self.after(
                3500,
                self.disarm
            )

            self.app.notify(
                "warning",
                (
                    f"Cliquez encore sur « Confirmer ? » "
                    f"pour supprimer « {phone['name']} »."
                ),
                3300
            )

            return

        # Deuxième étape
        self.disarm()

        try:

            self.app.stock.delete(
                phone["serial"]
            )

        except StockError as e:

            self.app.notify(
                "error",
                str(e),
                6000
            )

            return

        self.app.notify(
            "success",
            f"« {phone['name']} » a été supprimé."
        )

        self.refresh()

    def disarm(self):

        if self._arm_job:

            self.after_cancel(
                self._arm_job
            )

            self._arm_job = None

        if self.armed:

            self.armed = False

            self.btn_delete.configure(
                text="🗑  Supprimer",
                text_color=C["danger"]
            )

            self.btn_delete.set_colors(
                C["danger_soft"],
                C["danger_soft_hover"]
            )


# ============================================================
# 11. APPLICATION PRINCIPALE
# ============================================================

class App(ctk.CTk):

    NAV = (
        (
            "dashboard",
            "📊",
            "Dashboard"
        ),

        (
            "add",
            "➕",
            "Ajouter un téléphone"
        ),

        (
            "stock",
            "📦",
            "Stock"
        ),
    )

    def __init__(self):

        super().__init__()

        # ----------------------------------------------------
        # IMPORTANT :
        # Pas de self.attributes("-alpha", 0.0)
        # Pas de fade-in
        # ----------------------------------------------------

        self.title(
            f"{STORE_NAME} — {TAGLINE}"
        )

        self.configure(
            fg_color=C["bg"]
        )

        self._setup_window()

        self.stock = ExcelStock(
            os.path.join(
                app_dir(),
                EXCEL_FILENAME
            )
        )

        self.current = None

        self._anim_token = 0

        self.nav_buttons = {}

        self.grid_columnconfigure(
            1,
            weight=1
        )

        self.grid_rowconfigure(
            1,
            weight=1
        )

        # ----------------------------------------------------
        # Interface
        # ----------------------------------------------------

        self._build_topbar()

        self._build_sidebar()

        self.content = ctk.CTkFrame(
            self,
            fg_color=C["bg"],
            corner_radius=0
        )

        self.content.grid(
            row=1,
            column=1,
            sticky="nsew"
        )

        # ----------------------------------------------------
        # Pages
        # ----------------------------------------------------

        self.pages = {
            "dashboard": DashboardPage(self),
            "add": AddPage(self),
            "stock": StockPage(self)
        }

        self.banner = Banner(
            self.content
        )

        # ----------------------------------------------------
        # Première page
        # ----------------------------------------------------

        self.show_page(
            "dashboard",
            animate=False
        )

        # Horloge
        self._tick_clock()

    # ========================================================
    # FENÊTRE
    # ========================================================

    def _setup_window(self):

        screen_w = self.winfo_screenwidth()
        screen_h = self.winfo_screenheight()

        width = min(
            1280,
            int(screen_w * 0.92)
        )

        height = min(
            780,
            int(screen_h * 0.88)
        )

        x = max(
            0,
            (screen_w - width) // 2
        )

        y = max(
            0,
            (screen_h - height) // 2
        )

        self.geometry(
            f"{width}x{height}+{x}+{y}"
        )

        self.minsize(
            1000,
            620
        )

    # ========================================================
    # BARRE SUPÉRIEURE
    # ========================================================

    def _build_topbar(self):

        bar = ctk.CTkFrame(
            self,
            height=72,
            corner_radius=0,
            fg_color=C["card"]
        )

        bar.grid(
            row=0,
            column=0,
            columnspan=2,
            sticky="ew"
        )

        bar.pack_propagate(False)

        ctk.CTkFrame(
            bar,
            height=1,
            corner_radius=0,
            fg_color=C["border"]
        ).pack(
            side="bottom",
            fill="x"
        )

        # Logo
        ctk.CTkLabel(
            bar,
            text="📱",
            width=46,
            height=46,
            corner_radius=14,
            fg_color=C["accent"],
            text_color="white",
            font=F(22)
        ).pack(
            side="left",
            padx=(24, 12),
            pady=12
        )

        # Nom
        titles = ctk.CTkFrame(
            bar,
            fg_color="transparent"
        )

        titles.pack(
            side="left"
        )

        ctk.CTkLabel(
            titles,
            text=STORE_NAME,
            font=F(22, "bold"),
            text_color=C["text"],
            anchor="w"
        ).pack(
            anchor="w"
        )

        ctk.CTkLabel(
            titles,
            text=TAGLINE,
            font=F(12),
            text_color=C["muted"],
            anchor="w"
        ).pack(
            anchor="w"
        )

        # Heure
        right = ctk.CTkFrame(
            bar,
            fg_color="transparent"
        )

        right.pack(
            side="right",
            padx=26
        )

        self.clock_label = ctk.CTkLabel(
            right,
            text="",
            font=F(20, "bold"),
            text_color=C["text"],
            anchor="e"
        )

        self.clock_label.pack(
            anchor="e"
        )

        self.date_label = ctk.CTkLabel(
            right,
            text="",
            font=F(12),
            text_color=C["muted"],
            anchor="e"
        )

        self.date_label.pack(
            anchor="e"
        )

    # ========================================================
    # HORLOGE
    # ========================================================

    def _tick_clock(self):

        now = datetime.datetime.now()

        self.clock_label.configure(
            text=now.strftime("%H:%M")
        )

        self.date_label.configure(
            text=(
                f"{DAYS[now.weekday()].capitalize()} "
                f"{now.day} "
                f"{MONTHS[now.month - 1]} "
                f"{now.year}"
            )
        )

        self.after(
            1000,
            self._tick_clock
        )

    # ========================================================
    # MENU LATÉRAL
    # ========================================================

    def _build_sidebar(self):

        side = ctk.CTkFrame(
            self,
            width=236,
            corner_radius=0,
            fg_color=C["sidebar"]
        )

        side.grid(
            row=1,
            column=0,
            sticky="ns"
        )

        side.pack_propagate(False)

        ctk.CTkLabel(
            side,
            text="MENU",
            font=F(11, "bold"),
            text_color="#64748B",
            anchor="w"
        ).pack(
            fill="x",
            padx=26,
            pady=(26, 8)
        )

        for key, icon, label in self.NAV:

            btn = HoverButton(
                side,
                C["sidebar"],
                C["sidebar_hover"],
                text=f"   {icon}    {label}",
                anchor="w",
                height=48,
                corner_radius=12,
                font=F(14, "bold"),
                text_color=C["sidebar_text"],
                command=partial(
                    self.navigate,
                    key
                )
            )

            btn.pack(
                fill="x",
                padx=14,
                pady=3
            )

            self.nav_buttons[key] = btn

        ctk.CTkLabel(
            side,
            text=(
                f"📄  {EXCEL_FILENAME}\n"
                "Version 1.0"
            ),
            font=F(11),
            text_color="#64748B",
            justify="left",
            anchor="w"
        ).pack(
            side="bottom",
            fill="x",
            padx=26,
            pady=20
        )

    def _set_active(self, key):

        for name, btn in self.nav_buttons.items():

            if name == key:

                btn.set_colors(
                    C["accent"],
                    C["accent_hover"]
                )

                btn.configure(
                    text_color="white"
                )

            else:

                btn.set_colors(
                    C["sidebar"],
                    C["sidebar_hover"]
                )

                btn.configure(
                    text_color=C["sidebar_text"]
                )

    # ========================================================
    # NAVIGATION
    # ========================================================

    def navigate(self, key):

        if key == "add":

            self.pages["add"].set_add_mode()

        self.show_page(key)

    def start_edit(self, phone):

        self.pages["add"].load_for_edit(
            phone
        )

        self.show_page(
            "add"
        )

    def show_page(
        self,
        key,
        animate=True
    ):

        if key == self.current:

            self.pages[key].on_show()

            return

        self._anim_token += 1

        token = self._anim_token

        if self.current:

            self.pages[
                self.current
            ].place_forget()

        page = self.pages[key]

        self.current = key

        self._set_active(key)

        page.on_show()

        page.place(
            relx=0.05 if animate else 0,
            rely=0,
            relwidth=1,
            relheight=1
        )

        if animate:

            self._slide(
                page,
                token
            )

        if hasattr(
            self,
            "banner"
        ):

            self.banner.lift()

    def _slide(
        self,
        page,
        token,
        i=1,
        steps=14
    ):

        if token != self._anim_token:
            return

        page.place(
            relx=0.05 * (
                1 - ease_out(i / steps)
            ),
            rely=0,
            relwidth=1,
            relheight=1
        )

        if i < steps:

            self.after(
                14,
                lambda: self._slide(
                    page,
                    token,
                    i + 1,
                    steps
                )
            )

    # ========================================================
    # SERVICES
    # ========================================================

    def notify(
        self,
        kind,
        text,
        duration=3800
    ):

        self.banner.show(
            kind,
            text,
            duration
        )

    def fetch(self):

        try:

            return self.stock.get_all()

        except StockError as e:

            self.notify(
                "error",
                str(e),
                6000
            )

            return []


# ============================================================
# 12. LANCEMENT
# ============================================================

def main():

    ctk.set_appearance_mode(
        "light"
    )

    ctk.set_default_color_theme(
        "blue"
    )

    app = App()

    app.mainloop()


if __name__ == "__main__":
    main()